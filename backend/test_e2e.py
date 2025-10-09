import os
import time
import json
import requests

BACKEND = os.environ.get("BACKEND", "http://localhost:8000")

def ensure_sample_docx(path: str):
    from docx import Document
    d = Document()
    d.add_heading('Sample Document', 0)
    d.add_paragraph('The ACME Product is a fictional product for testing. It features fast performance and low cost.')
    d.add_paragraph('Use cases include analytics and reporting. The warranty period is 1 year.')
    d.save(path)


def main():
    sample_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'examples', 'sample.docx'))
    os.makedirs(os.path.dirname(sample_path), exist_ok=True)
    ensure_sample_docx(sample_path)

    files = { 'file': open(sample_path, 'rb') }
    data = {}
    print('Uploading sample.docx ...')
    r = requests.post(f"{BACKEND}/upload", files=files, data=data, timeout=120)
    print('Upload status:', r.status_code)
    print('Upload body:', r.text)
    r.raise_for_status()
    up = r.json()
    session_id = up.get('session_id')
    if not session_id:
        raise SystemExit('No session_id in upload response')

    time.sleep(1)
    q = "What is ACME Product and what is the warranty?"
    print('Asking:', q)
    r2 = requests.post(f"{BACKEND}/ask", json={
        'session_id': session_id,
        'question': q,
        'k': 5,
        'gemini_api_key': None
    }, timeout=120)
    print('Ask status:', r2.status_code)
    print('Ask body:', r2.text)
    r2.raise_for_status()
    ans = r2.json()
    print('\nAnswer:')
    print(ans.get('answer'))

if __name__ == '__main__':
    main()
